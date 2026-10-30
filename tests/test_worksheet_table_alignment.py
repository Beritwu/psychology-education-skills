import pytest

from test_benchmark_evaluation import client  # noqa: F401


def test_canonical_record_task_without_fields_exports_writable_student_table(client):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server
    import io
    import docx

    class AssessmentTableLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'worksheet':
                result['tasks'] = [{
                    'task_id': 'quiz', 'task_role': 'assessment', 'kind': 'record_table',
                    'prompt': '在作答记录表中回答评价题。', 'fields': [],
                    'assessment_refs': ['a1'],
                }]
            return result

    base = '/api/projects/assessment-record-table'
    client.post('/api/projects', json={'project_id': 'assessment-record-table'})
    engine = server.get_engine('assessment-record-table')
    engine.llm, engine.research = AssessmentTableLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    rev = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': rev}).json()['content']
    files = {f['kind']: f for f in manifest['files'] if f['role'] == 'worksheet'}
    text = client.get(base + '/file', params={'path': files['markdown']['path']}).text
    assert '| 作答要点 | 记录 |' in text
    assert text.count('|  |  |') >= 4
    assert '**答案：**' not in text
    response = client.get(base + '/file', params={'path': files['docx']['path']})
    document = docx.Document(io.BytesIO(response.content))
    table = next(table for table in document.tables if table.cell(0, 0).text == '作答要点')
    assert len(table.rows) >= 5
    from docx.shared import Cm
    assert table.rows[1].height >= Cm(0.6)
    assert all('**' not in paragraph.text for paragraph in document.paragraphs)


def test_export_table_adds_missing_label_column_for_student_and_teacher(client):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server

    class TableLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'worksheet':
                result['tasks'] = [{
                    'task_id': 'matrix', 'kind': 'record_table', 'prompt': '分类',
                    'fields': ['刺激变化', '行为变化', '术语'], 'row_labels': ['情形一'],
                    'answer': {'情形一': {'刺激变化': '加入', '行为变化': '增加', '术语': '正强化'}},
                }]
            return result

    project = 'table-alignment'
    base = '/api/projects/' + project
    client.post('/api/projects', json={'project_id': project})
    engine = server.get_engine(project)
    engine.llm, engine.research = TableLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    revision = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': revision}).json()['content']
    files = [f for f in manifest['files'] if f['role'] in ('worksheet', 'worksheet_teacher') and f['kind'] == 'markdown']
    assert len(files) == 2
    outputs = [client.get(base + '/file', params={'path': f['path']}).text for f in files]
    for output in outputs:
        assert '| 条目 | 刺激变化 | 行为变化 | 术语 |' in output
        assert '| 情形一 |  |  |  |' in output
    assert any('| 情形一 | 加入 | 增加 | 正强化 |' in output for output in outputs)


@pytest.mark.parametrize('answer_form', ['list', 'mapping'])
def test_export_unlabelled_table_provides_all_blank_rows_and_aligned_teacher_answers(client, answer_form):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server

    class FourRowsLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'worksheet':
                result['tasks'] = [{
                    'task_id': 'matrix', 'kind': 'record_table', 'prompt': '填写四种组合',
                    'fields': ['刺激变化', '行为变化', '术语'],
                    'answer': [
                        {'row_label': '组合一', '刺激变化': '加入', '行为变化': '增加', '术语': '正强化'},
                        {'row_label': '组合二', '刺激变化': '移除', '行为变化': '增加', '术语': '负强化'},
                        {'row_label': '组合三', '刺激变化': '加入', '行为变化': '减少', '术语': '正惩罚'},
                        {'row_label': '组合四', '刺激变化': '移除', '行为变化': '减少', '术语': '负惩罚'},
                    ],
                }]
                if answer_form == 'mapping':
                    rows = result['tasks'][0]['answer']
                    result['tasks'][0]['answer'] = {
                        row['row_label']: {key: value for key, value in row.items() if key != 'row_label'}
                        for row in rows
                    }
            return result

    project = 'table-unlabelled-rows-' + answer_form
    base = '/api/projects/' + project
    client.post('/api/projects', json={'project_id': project})
    engine = server.get_engine(project)
    engine.llm, engine.research = FourRowsLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    revision = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': revision}).json()['content']
    files = {f['role']: f for f in manifest['files'] if f['role'] in ('worksheet', 'worksheet_teacher') and f['kind'] == 'markdown'}
    student = client.get(base + '/file', params={'path': files['worksheet']['path']}).text
    teacher = client.get(base + '/file', params={'path': files['worksheet_teacher']['path']}).text
    assert student.count('|  |  |  |') >= 4
    assert '正强化' not in student
    assert '| 条目 | 刺激变化 | 行为变化 | 术语 |' in teacher
    assert '| 组合一 | 加入 | 增加 | 正强化 |' in teacher


def test_export_hides_answer_bearing_question_identifiers_in_both_formats(client):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server
    import io
    import docx

    secret_id = 'sec1_hinge_pr'

    class LabelLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'assessment':
                result['items'][0]['assessment_id'] = secret_id
                result['items'][0]['remediation'] = {'A': '回看刺激变化', 'C': '重新比较行为'}
                result['items'][0]['max_score'] = 4
                result['items'][0]['scoring'] = {'total_points': 4, 'criteria': [
                    {'criterion': '说明证据', 'points': 4, 'partial_credit': '部分正确得2分'}]}
            if kwargs['task'] == 'worksheet':
                result['tasks'] = [{
                    'task_id': 'quiz', 'task_role': 'assessment', 'kind': 'record_table',
                    'prompt': '在表中完成 ' + secret_id, 'fields': ['题目', '作答'],
                    'row_labels': [secret_id], 'assessment_refs': [secret_id],
                    'answer': '仅供教师',
                }]
            return result

    project = 'question-label-privacy'
    base = '/api/projects/' + project
    client.post('/api/projects', json={'project_id': project})
    engine = server.get_engine(project)
    engine.llm, engine.research = LabelLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    revision = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': revision}).json()['content']
    files = [f for f in manifest['files'] if f['role'] in ('worksheet', 'worksheet_teacher') and f['kind'] in ('markdown', 'docx')]
    assert len(files) == 3
    for f in files:
        response = client.get(base + '/file', params={'path': f['path']})
        if f['kind'] == 'docx':
            document = docx.Document(io.BytesIO(response.content))
            text = '\n'.join(p.text for p in document.paragraphs)
            text += '\n' + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            text = response.text
        assert secret_id not in text
        assert '主评价 1' in text
        if f['kind'] == 'markdown':
            assert '| 主评价 1 |' in text
        else:
            assert any(cell.text == '主评价 1' for table in document.tables
                       for row in table.rows for cell in row.cells)
        if f['role'] == 'worksheet_teacher':
            assert '- A：回看刺激变化' in text
            assert "{'A':" not in text
            assert '说明证据' in text and '部分正确得2分' in text and '4 分' in text
        else:
            assert '部分正确得2分' not in text


def test_short_question_id_does_not_rewrite_dates_or_options(client):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server

    class ShortIdLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'assessment':
                result['items'][0]['assessment_id'] = '1'
                result['items'][0]['question'] = '日期2024-1-1，选择 A 或 B。'
            if kwargs['task'] == 'worksheet':
                result['tasks'] = [{
                    'task_id': 'quiz', 'task_role': 'assessment', 'kind': 'record_table',
                    'prompt': '请作答', 'fields': ['题目', '作答'],
                    'row_labels': ['1'], 'assessment_refs': ['1'], 'answer': '仅供教师',
                }]
            return result

    base = '/api/projects/question-id-short'
    client.post('/api/projects', json={'project_id': 'question-id-short'})
    engine = server.get_engine('question-id-short')
    engine.llm, engine.research = ShortIdLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    rev = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': rev}).json()['content']
    student_file = next(f for f in manifest['files'] if f['role'] == 'worksheet' and f['kind'] == 'markdown')
    student = client.get(base + '/file', params={'path': student_file['path']}).text
    assert '日期2024-1-1，选择 A 或 B。' in student
    assert '### 主评价 1' in student
    assert '| 主评价 1 |' in student


def test_export_hides_internal_question_id_in_student_script(client):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server
    import io
    import docx

    secret_id = 'sec1_hinge_pr'
    class ScriptLabelLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'assessment':
                result['items'][0]['assessment_id'] = secret_id
            if kwargs['task'] == 'worksheet':
                result['tasks'][0].update(task_role='assessment', assessment_refs=[secret_id])
            if kwargs['task'] in ('script', 'script_fix'):
                result['units'][1]['text'] += ' 请完成题目' + secret_id + '。'
            return result

    base = '/api/projects/script-label-privacy'
    client.post('/api/projects', json={'project_id': 'script-label-privacy'})
    engine = server.get_engine('script-label-privacy')
    engine.llm, engine.research = ScriptLabelLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    rev = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': rev}).json()['content']
    files = [f for f in manifest['files'] if f['role'] == 'lesson_script' and f['kind'] in ('markdown', 'docx')]
    assert len(files) == 2
    for f in files:
        response = client.get(base + '/file', params={'path': f['path']})
        if f['kind'] == 'docx':
            document = docx.Document(io.BytesIO(response.content))
            text = '\n'.join(p.text for p in document.paragraphs)
            text += '\n' + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            text = response.text
        assert secret_id not in text
        assert '主评价 1' in text


def test_export_shows_verification_rubric_only_to_teacher(client):
    from conftest import FakeLLM, FakeResearch, run_build
    from pebs import server
    import io
    import docx

    rubric = '评分口径：路径各2分，时间顺序1分，总分5分。'
    class RubricLLM(FakeLLM):
        def generate_json(self, **kwargs):
            result = super().generate_json(**kwargs)
            if kwargs['task'] == 'assessment':
                result['items'][0]['answer_verification'] = rubric
            if kwargs['task'] == 'worksheet':
                result['tasks'] = [{'task_id': 'quiz', 'kind': 'record_table', 'task_role': 'assessment',
                                    'prompt': '完成题目', 'fields': ['题目', '作答'],
                                    'assessment_refs': [result_ref], 'answer': '教师专用'}]
            return result
    result_ref = 'a1'
    project = 'verification-rubric-export'
    base = '/api/projects/' + project
    client.post('/api/projects', json={'project_id': project})
    engine = server.get_engine(project)
    engine.llm, engine.research = RubricLLM(), FakeResearch()
    run_build(engine, '生成学习单与讲稿，不要PPT，不要动画。', environment='test_fixture')
    revision = engine.store.revisions_of('export_manifest')[-1]
    manifest = client.get(base + '/artifacts/export_manifest', params={'rev': revision}).json()['content']
    files = [f for f in manifest['files'] if f['role'] in ('worksheet', 'worksheet_teacher') and f['kind'] in ('markdown', 'docx')]
    assert len(files) == 3
    for f in files:
        response = client.get(base + '/file', params={'path': f['path']})
        if f['kind'] == 'docx':
            document = docx.Document(io.BytesIO(response.content))
            text = '\n'.join(p.text for p in document.paragraphs)
            text += '\n' + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            text = response.content.decode("utf-8")
        assert (rubric in text) == (f['role'] == 'worksheet_teacher')
